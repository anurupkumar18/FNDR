"""The readout-docx skill's renderer, on a fixture spec.

Needs python-docx and matplotlib, which FNDR installs into the hand-off
toolchain; skipped where they are absent (CI):
  python3 -m unittest scripts/hand-off/test_render_readout.py
"""

import importlib.util
import tempfile
import unittest
from pathlib import Path

RENDER = Path(__file__).resolve().parents[2] / "src-tauri/sidecars/hand-off/skills/readout-docx/render.py"
HAVE_LIBS = all(importlib.util.find_spec(m) for m in ("docx", "matplotlib"))

SPEC = {
    "title": "Q4 onboarding readout",
    "subtitle": "400 respondents",
    "summary": "Setup time drives churn.",
    "sections": [{"heading": "Drivers", "paragraphs": ["Setup comes first."], "bullets": ["Fix setup"]}],
    "table": {"columns": ["Rank", "Driver", "Churned", "Share"],
              "rows": [["1", "Setup took too long", "24", "30.8%"], ["2", "Missing Salesforce integration", "21", "26.9%"]]},
    "chart": {"title": "Churn drivers", "labels": ["Setup", "Salesforce"], "values": [24, 21], "value_label": "Churned"},
    "sources_note": "Data: inputs/survey.csv",
}


def load():
    spec = importlib.util.spec_from_file_location("render", RENDER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@unittest.skipUnless(HAVE_LIBS, "python-docx and matplotlib are in FNDR's hand-off toolchain, not here")
class RenderReadoutTests(unittest.TestCase):
    def test_a_spec_becomes_a_one_page_document_with_its_table_and_chart(self):
        import docx

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out" / "readout.docx"
            out.parent.mkdir()
            load().render(SPEC, out)
            document = docx.Document(out)
            text = [p.text for p in document.paragraphs if p.text]
            self.assertEqual(text[0], "Q4 onboarding readout")
            self.assertIn("Setup time drives churn.", text)
            rows = [[c.text for c in r.cells] for r in document.tables[0].rows]
            self.assertEqual(rows[0], ["Rank", "Driver", "Churned", "Share"])
            self.assertEqual(rows[2][1], "Missing Salesforce integration")
            self.assertTrue(document.tables[0].rows[0].cells[0].paragraphs[0].runs[0].bold)
            self.assertEqual(len(document.inline_shapes), 1)
            self.assertTrue((out.parent / "readout-chart.png").read_bytes().startswith(b"\x89PNG"))

    def test_a_sections_table_sits_under_that_section(self):
        import docx

        spec = {
            "title": "T",
            "summary": "S",
            "sections": [
                {"heading": "Drivers", "paragraphs": ["Ranked."], "table": SPEC["table"]},
                {"heading": "Next", "bullets": ["Fix setup"]},
            ],
        }
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out" / "readout.docx"
            out.parent.mkdir()
            load().render(spec, out)
            body = docx.Document(out).element.body
            order = [child.tag.split("}")[1] for child in body]
            texts = ["".join(child.itertext()) for child in body]
            at = lambda text: next(i for i, t in enumerate(texts) if text in t)
            table_at = order.index("tbl")
            self.assertLess(at("Ranked."), table_at)
            self.assertGreater(at("Next"), table_at)

    def test_only_a_docx_inside_out_is_written(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                load().render(SPEC, Path(tmp) / "readout.docx")
            with self.assertRaises(ValueError):
                load().render(SPEC, Path(tmp) / "out" / "readout.pdf")


if __name__ == "__main__":
    unittest.main()
