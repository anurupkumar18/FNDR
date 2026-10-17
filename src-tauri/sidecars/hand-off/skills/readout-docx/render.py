"""Render a one-page readout .docx (and chart .png) from a JSON spec."""
import json, sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Helvetica Neue"
INK = RGBColor(0x1F, 0x2A, 0x37)
MUTED = RGBColor(0x5B, 0x65, 0x70)
ACCENT = "3B6EA5"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def pad_cells(table, points=5):
    """Cell margins on every side, so text does not touch the borders."""
    props = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), str(points * 20))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    props.append(margins)


def is_number(text):
    return text.replace(",", "").replace("%", "").replace(".", "", 1).lstrip("-").isdigit()

def render(spec, out):
    """Write `out` (a .docx inside an `out/` folder) and, with a chart, its .png."""
    out = Path(out)
    if out.suffix != ".docx" or out.parent.name != "out":
        raise ValueError("The output must be a .docx inside out/")

    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin = section.bottom_margin = Inches(0.7)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.font.color.rgb = INK
    for name in ("Title", "Heading 2"):
        heading = doc.styles[name]
        heading.font.name = FONT
        heading.font.color.rgb = INK
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 2"].font.size = Pt(12.5)

    title = doc.add_heading(spec["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if spec.get("subtitle"):
        p = doc.add_paragraph(spec["subtitle"])
        p.runs[0].font.color.rgb = MUTED
    summary = doc.add_paragraph()
    summary.add_run(spec["summary"]).bold = True

    for sec in spec.get("sections", [])[:3]:
        doc.add_heading(sec["heading"], level=2)
        for para in sec.get("paragraphs", []):
            doc.add_paragraph(para)
        for bullet in sec.get("bullets", []):
            doc.add_paragraph(bullet, style="List Bullet")

    table = spec.get("table")
    if table:
        columns = [str(c) for c in table["columns"]]
        rows = [[str(v) for v in row] for row in table["rows"][:6]]
        t = doc.add_table(rows=1, cols=len(columns))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        pad_cells(t)
        numeric = [bool(rows) and all(is_number(r[i]) for r in rows if i < len(r)) for i in range(len(columns))]
        # A number column is as wide as its header needs; text columns share the rest.
        widths = [Inches(max(0.6, 0.11 * len(columns[i]) + 0.3)) if numeric[i] else None for i in range(len(columns))]
        free = Inches(6.9) - sum(w for w in widths if w)
        text_columns = sum(1 for w in widths if w is None) or 1
        widths = [w or int(free / text_columns) for w in widths]
        for i, (cell, name) in enumerate(zip(t.rows[0].cells, columns)):
            cell.text = name
            run = cell.paragraphs[0].runs[0]
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade(cell, ACCENT)
        for row in rows:
            t.add_row()
        for i, column in enumerate(t.columns):
            column.width = widths[i]
        for r, row in enumerate(t.rows):
            for i, cell in enumerate(row.cells):
                cell.width = widths[i]
                if r > 0 and i < len(rows[r - 1]):
                    cell.text = rows[r - 1][i]
                if numeric[i]:
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if r > 0 and r % 2 == 0:
                    shade(cell, "EEF3F8")
        doc.add_paragraph()

    chart = spec.get("chart")
    if chart:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        labels, values = chart["labels"][:6], chart["values"][:6]
        fig, ax = plt.subplots(figsize=(6.2, 2.4), dpi=200)
        bars = ax.barh(labels[::-1], values[::-1], color="#3b6ea5")
        ax.bar_label(bars, padding=3, fontsize=8)
        ax.set_xlabel(chart.get("value_label", ""), fontsize=8)
        ax.set_title(chart.get("title", ""), fontsize=9, loc="left")
        ax.spines[["top", "right"]].set_visible(False)
        ax.tick_params(labelsize=8)
        fig.tight_layout()
        png = out.with_name(out.stem + "-chart.png")
        fig.savefig(png)
        doc.add_picture(str(png), width=Inches(6.2))

    if spec.get("sources_note"):
        p = doc.add_paragraph(spec["sources_note"])
        p.runs[0].font.size = Pt(8)
        p.runs[0].font.color.rgb = MUTED

    doc.save(out)


if __name__ == "__main__":
    try:
        render(json.loads(Path(sys.argv[1]).read_text()), sys.argv[2])
    except ValueError as error:
        sys.exit(str(error))
    print(sys.argv[2])
